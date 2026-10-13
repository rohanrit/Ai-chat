"""Fully local, cross-platform file content extraction.

No online APIs, no system-level binaries (no Tesseract/poppler) are required.
OCR runs through RapidOCR (ONNX, models bundled in the wheel) and scanned PDF
pages are rasterized with PyMuPDF.

Extraction supports:
  * live progress reporting via ``progress_callback(percent, stage)``
  * cooperative cancellation via a ``threading.Event`` (``stop_event``)
  * a hard input file-size limit (``MAX_FILE_SIZE_MB``)
"""

import os
from pathlib import Path

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}
TEXT_EXTENSIONS = {
    ".txt", ".py", ".md", ".json", ".csv", ".log", ".xml", ".yaml", ".yml",
    ".ini", ".cfg", ".toml", ".html", ".css", ".js", ".ts", ".java", ".c",
    ".cpp", ".h", ".sh", ".bat", ".ps1", ".sql", ".rs", ".go",
}

# Hard cap on the raw input file size (bytes) that we will attempt to read.
MAX_FILE_SIZE_MB = int(os.getenv("MAX_FILE_SIZE_MB", "25"))
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

_OCR_ENGINE = None


class ScanCancelled(Exception):
    """Raised when the user aborts a running file scan."""


class FileTooLargeError(Exception):
    """Raised when an input file exceeds the supported processing limit."""


def get_ocr():
    """Return a cached RapidOCR engine (models load once, ~0.5s)."""
    global _OCR_ENGINE
    if _OCR_ENGINE is None:
        from rapidocr_onnxruntime import RapidOCR
        _OCR_ENGINE = RapidOCR()
    return _OCR_ENGINE


def _check_stop(stop_event):
    if stop_event is not None and stop_event.is_set():
        raise ScanCancelled()


def _report(progress_callback, percent, stage):
    if progress_callback:
        progress_callback(max(0.0, min(100.0, float(percent))), stage)


def _ocr_image(image):
    """Run local OCR on a PIL image and return the joined detected text."""
    import numpy as np

    if image.mode != "RGB":
        image = image.convert("RGB")
    array = np.array(image)
    result, _ = get_ocr()(array)
    if not result:
        return ""
    return "\n".join(line[1] for line in result if len(line) > 1 and line[1]).strip()


def _extract_image(path, progress_callback=None, stop_event=None):
    from PIL import Image

    _report(progress_callback, 5, "opening image")
    _check_stop(stop_event)
    
    file_size_bytes = path.stat().st_size
    file_size_kb = file_size_bytes / 1024

    with Image.open(path) as img:
        width, height = img.size
        img_format = img.format or path.suffix.upper().lstrip(".")
        img_mode = img.mode
        
        _report(progress_callback, 30, "running local OCR")
        _check_stop(stop_event)
        ocr_text = _ocr_image(img)

    header = (
        f"[Image Metadata]\n"
        f"File Name: {path.name}\n"
        f"Dimensions: {width} x {height} pixels\n"
        f"File Size: {file_size_kb:.1f} KB ({file_size_bytes} bytes)\n"
        f"Format: {img_format} | Mode: {img_mode}\n\n"
        f"[OCR Text Content]\n"
        f"{ocr_text if ocr_text else '(No text detected via OCR)'}"
    )

    _report(progress_callback, 100, "done")
    return "image", header


def _extract_pdf(path, progress_callback=None, stop_event=None):
    import pypdf

    reader = pypdf.PdfReader(str(path))
    total = max(1, len(reader.pages))
    parts = []
    text_pages = 0
    ocr_pages = 0

    doc = None
    for i, page in enumerate(reader.pages):
        _check_stop(stop_event)
        _report(progress_callback, (i / total) * 100, f"page {i + 1}/{total}")

        page_text = (page.extract_text() or "").strip()
        if len(page_text) >= 20:
            parts.append(f"--- Page {i + 1} (text) ---\n{page_text}")
            text_pages += 1
            _report(progress_callback, ((i + 1) / total) * 100, f"page {i + 1}/{total}")
            continue

        try:
            if doc is None:
                import pymupdf
                doc = pymupdf.open(str(path))
            pix = doc[i].get_pixmap(dpi=200)
            from PIL import Image

            image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            _check_stop(stop_event)
            ocr_text = _ocr_image(image)
        except ScanCancelled:
            raise
        except Exception:
            ocr_text = ""

        if ocr_text:
            parts.append(f"--- Page {i + 1} (ocr) ---\n{ocr_text}")
            ocr_pages += 1
        elif page_text:
            parts.append(f"--- Page {i + 1} ---\n{page_text}")
            text_pages += 1

        _report(progress_callback, ((i + 1) / total) * 100, f"page {i + 1}/{total}")

    if doc is not None:
        doc.close()

    _report(progress_callback, 100, "done")
    return "pdf", "\n\n".join(parts).strip(), text_pages, ocr_pages


def _extract_docx(path, progress_callback=None, stop_event=None):
    import docx

    _report(progress_callback, 5, "opening document")
    _check_stop(stop_event)
    doc = docx.Document(str(path))
    blocks = []

    paragraphs = doc.paragraphs
    total_paras = max(1, len(paragraphs))
    for idx, para in enumerate(paragraphs):
        _check_stop(stop_event)
        if para.text.strip():
            blocks.append(para.text)
        if idx % 25 == 0:
            _report(progress_callback, (idx / total_paras) * 90, "reading paragraphs")

    for table in doc.tables:
        _check_stop(stop_event)
        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells if cell.text.strip()
            )
            if row_text:
                blocks.append(row_text)

    _report(progress_callback, 100, "done")
    return "docx", "\n".join(blocks).strip()


def _extract_text(path, progress_callback=None, stop_event=None):
    size = 0
    try:
        size = path.stat().st_size
    except OSError:
        size = 0

    _report(progress_callback, 5, "reading file")

    def _read(encoding, errors=None):
        chunks = []
        read = 0
        with open(path, "r", encoding=encoding, errors=errors) as handle:
            while True:
                _check_stop(stop_event)
                chunk = handle.read(65536)
                if not chunk:
                    break
                chunks.append(chunk)
                read += len(chunk)
                if size:
                    _report(progress_callback, min(99, (read / size) * 100), "reading file")
        return "".join(chunks)

    try:
        text = _read("utf-8")
    except UnicodeDecodeError:
        text = _read("latin-1", errors="ignore")

    _report(progress_callback, 100, "done")
    return "text", text


def extract_content(path, progress_callback=None, stop_event=None):
    """Extract readable text from a local file.

    Args:
        path: file path to read.
        progress_callback: optional ``callable(percent, stage)`` for live UI.
        stop_event: optional ``threading.Event``; when set, extraction aborts
            with :class:`ScanCancelled`.

    Returns a dict ``{"kind", "text", "meta"}``. Raises
    :class:`FileTooLargeError`, :class:`ScanCancelled`, or ``ValueError``.
    """
    path = Path(path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")

    try:
        size_bytes = path.stat().st_size
    except OSError:
        size_bytes = 0

    if size_bytes > MAX_FILE_SIZE_BYTES:
        size_mb = size_bytes / (1024 * 1024)
        raise FileTooLargeError(
            f"'{path.name}' is {size_mb:.1f} MB, which exceeds the maximum "
            f"supported size of {MAX_FILE_SIZE_MB} MB."
        )

    ext = path.suffix.lower()
    try:
        if ext in IMAGE_EXTENSIONS:
            kind, text = _extract_image(path, progress_callback, stop_event)
            return {"kind": kind, "text": text, "meta": {"ocr": True}}

        if ext == ".pdf":
            kind, text, text_pages, ocr_pages = _extract_pdf(path, progress_callback, stop_event)
            return {
                "kind": kind,
                "text": text,
                "meta": {"text_pages": text_pages, "ocr_pages": ocr_pages},
            }

        if ext == ".docx":
            kind, text = _extract_docx(path, progress_callback, stop_event)
            return {"kind": kind, "text": text, "meta": {}}

        if ext in TEXT_EXTENSIONS or not ext:
            kind, text = _extract_text(path, progress_callback, stop_event)
            return {"kind": kind, "text": text, "meta": {}}

        # Unknown extension: attempt a best-effort text read.
        try:
            kind, text = _extract_text(path, progress_callback, stop_event)
            return {"kind": kind, "text": text, "meta": {}}
        except ScanCancelled:
            raise
        except Exception as exc:
            raise ValueError(f"Unsupported or unreadable file type: {ext}") from exc
    except ScanCancelled:
        raise
    except FileTooLargeError:
        raise
